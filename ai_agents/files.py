# Upload validation, storage and text extraction for agent documents.
# Files are never executed or served inline: they are stored under a random name
# outside /static, text is extracted once, and only that text reaches the model.

import hashlib
import io
import os
import uuid
import zipfile
from dataclasses import dataclass
from typing import List

from flask import current_app
from werkzeug.datastructures import FileStorage
from werkzeug.utils import secure_filename

from extensions import db

from .models import AgentFile

TEXT_EXTENSIONS = {".txt", ".md", ".csv"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {".pdf", ".docx"}

CONTENT_TYPES = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".csv": "text/csv",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}

# Zip-bomb guard for .docx
MAX_DOCX_UNCOMPRESSED = 50 * 1024 * 1024
MAX_PDF_PAGES = 300


class FileValidationError(ValueError):
    pass


@dataclass
class ExtractedText:
    text: str
    truncated: bool


def upload_dir(company_id: int) -> str:
    base = current_app.config["AI_UPLOAD_DIR"]
    path = os.path.join(base, str(int(company_id)))
    os.makedirs(path, exist_ok=True)
    return path


def stored_path(agent_file: AgentFile) -> str:
    return os.path.join(upload_dir(agent_file.company_id), agent_file.stored_name)


def save_uploads(storages: List[FileStorage], *, company_id: int, user_id: int,
                 allowed_extensions) -> List[AgentFile]:
    """Validate every file first, then store them. Raises FileValidationError on the first bad file."""
    max_bytes = current_app.config["AI_MAX_FILE_BYTES"]
    max_chars = current_app.config["AI_MAX_DOCUMENT_CHARS"]

    prepared = []
    for storage in storages:
        original = (storage.filename or "").strip()
        safe = secure_filename(original)
        ext = os.path.splitext(safe)[1].lower()
        label = original or "file"
        if not safe or not ext:
            raise FileValidationError(f"'{label}' has no usable file name or extension.")
        if ext not in allowed_extensions or ext not in SUPPORTED_EXTENSIONS:
            allowed = ", ".join(sorted(allowed_extensions))
            raise FileValidationError(f"'{label}' is not a supported file type. Allowed: {allowed}.")

        data = storage.read(max_bytes + 1)
        if len(data) > max_bytes:
            raise FileValidationError(f"'{label}' is larger than the {max_bytes // (1024 * 1024)} MB limit.")
        if not data:
            raise FileValidationError(f"'{label}' is empty.")

        extracted = extract_text(data, ext, label, max_chars)
        prepared.append((original[:255], ext, data, extracted))

    saved = []
    for original, ext, data, extracted in prepared:
        stored_name = f"{uuid.uuid4().hex}{ext}"
        record = AgentFile(
            company_id=company_id,
            uploaded_by_id=user_id,
            original_filename=original,
            stored_name=stored_name,
            extension=ext,
            content_type=CONTENT_TYPES[ext],
            size_bytes=len(data),
            sha256=hashlib.sha256(data).hexdigest(),
            extracted_text=extracted.text,
            extracted_chars=len(extracted.text),
            truncated=extracted.truncated,
        )
        with open(os.path.join(upload_dir(company_id), stored_name), "wb") as fh:
            fh.write(data)
        db.session.add(record)
        saved.append(record)
    return saved


def extract_text(data: bytes, ext: str, label: str, max_chars: int) -> ExtractedText:
    if ext in TEXT_EXTENSIONS:
        text = _decode_text(data, label)
    elif ext == ".pdf":
        text = _pdf_text(data, label)
    elif ext == ".docx":
        text = _docx_text(data, label)
    else:  # pragma: no cover - guarded by SUPPORTED_EXTENSIONS
        raise FileValidationError(f"'{label}' is not a supported file type.")

    text = text.strip()
    if not text:
        raise FileValidationError(f"No readable text was found in '{label}'. Scanned PDFs are not supported yet.")
    if len(text) > max_chars:
        return ExtractedText(text=text[:max_chars], truncated=True)
    return ExtractedText(text=text, truncated=False)


def _decode_text(data: bytes, label: str) -> str:
    if b"\x00" in data:
        raise FileValidationError(f"'{label}' does not look like a text file.")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise FileValidationError(f"'{label}' is not in a readable text encoding (use UTF-8).")


def _pdf_text(data: bytes, label: str) -> str:
    if not data.startswith(b"%PDF-"):
        raise FileValidationError(f"'{label}' is not a valid PDF.")
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise FileValidationError("PDF support is not installed on this server (pypdf).") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise FileValidationError(f"'{label}' is password protected.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise FileValidationError(f"'{label}' has more than {MAX_PDF_PAGES} pages.")
        pages = []
        for number, page in enumerate(reader.pages, start=1):
            pages.append(f"[Page {number}]\n{page.extract_text() or ''}")
        return "\n\n".join(pages)
    except FileValidationError:
        raise
    except Exception as exc:
        raise FileValidationError(f"'{label}' could not be read as a PDF.") from exc


def _docx_text(data: bytes, label: str) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = zf.namelist()
            if "word/document.xml" not in names:
                raise FileValidationError(f"'{label}' is not a valid Word document.")
            if sum(info.file_size for info in zf.infolist()) > MAX_DOCX_UNCOMPRESSED:
                raise FileValidationError(f"'{label}' is too large once decompressed.")
    except zipfile.BadZipFile as exc:
        raise FileValidationError(f"'{label}' is not a valid Word document.") from exc

    try:
        import docx
    except ImportError as exc:
        raise FileValidationError("Word document support is not installed on this server (python-docx).") from exc
    try:
        document = docx.Document(io.BytesIO(data))
        lines = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
        return "\n".join(lines)
    except Exception as exc:
        raise FileValidationError(f"'{label}' could not be read as a Word document.") from exc
